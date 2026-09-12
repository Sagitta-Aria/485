"""Correlated external-short batch calibration and report generation.

The GUI owns TCP routing. This module owns the measurement schedule, waits for
matching RESULT frames, solves one residual path resistance per matrix port,
and writes an archival DOCX plus a machine-readable calibration candidate. The
fixture must externally short all tested ports. This module does not apply the
candidate values to normal cable measurements.
"""

from __future__ import annotations

from cable_tester.paths import REQUIREMENTS, CALIBRATION_REPORTS

import json
import queue
import re
import statistics
import sys
import threading
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable


DEFAULT_REPEAT_COUNT = 3
DEFAULT_SETTLE_SECONDS = 2.0
DEFAULT_COOLDOWN_SECONDS = 0.2
DEFAULT_RESPONSE_TIMEOUT_SECONDS = 5.0
PORTS_PER_BANK = 24
REPORT_ROOT = CALIBRATION_REPORTS
MEASURE_OK_PATTERN = re.compile(
    r"^OK MEASURE resistance=(?P<resistance>\d+(?:\.\d+)?) "
    r"raw=(?P<raw>\d+) range=(?P<range>\d+)$"
)

SendRequest = Callable[[str, str, str], str]
PublishEvent = Callable[[str, object], None]
ReportWriter = Callable[["BatchSession", Path], tuple[Path, Path]]


def ensure_report_dependency() -> None:
    """Fail before measurement when this interpreter cannot generate DOCX."""
    try:
        __import__("docx")
    except ImportError as error:
        requirements = REQUIREMENTS
        raise RuntimeError(
            "当前运行上位机的Python缺少python-docx，请执行："
            f'"{sys.executable}" -m pip install -r "{requirements}"'
        ) from error


@dataclass(frozen=True)
class MatrixPort:
    """Identify one CH446 X pin within the S1 or S2 bank."""

    bank: str
    x: int

    def __post_init__(self) -> None:
        if self.bank not in {"S1", "S2"} or not 0 <= self.x < PORTS_PER_BANK:
            raise ValueError(f"invalid matrix port: {self.bank} X{self.x}")

    @property
    def label(self) -> str:
        """Return the stable label used in reports and JSON output."""
        return f"{self.bank}_X{self.x}"


@dataclass(frozen=True)
class PairSpec:
    """Describe one pair equation or a cross-bank validation measurement."""

    positive: MatrixPort
    negative: MatrixPort
    purpose: str

    @property
    def key(self) -> str:
        """Return an order-preserving unique key for result aggregation."""
        return f"{self.positive.label}__{self.negative.label}"

    @property
    def command(self) -> str:
        """Return the normal four-wire routing command for this pair."""
        return (
            f"CONNECT {self.positive.bank} {self.positive.x} "
            f"{self.negative.bank} {self.negative.x}"
        )

    @property
    def label(self) -> str:
        """Return a compact human-readable pair label."""
        return f"{self.positive.label} - {self.negative.label}"


@dataclass(frozen=True)
class CalibrationReading:
    """Hold one successfully decoded normal MEASURE response."""

    resistance_ohm: float
    raw_value: int
    range_code: int
    payload: str
    received_at: str


@dataclass
class PairSamples:
    """Accumulate successful readings and exact error payloads for one pair."""

    spec: PairSpec
    readings: list[CalibrationReading] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def representative_ohm(self) -> float | None:
        """Use the median so one noisy reading does not dominate the equation."""
        if not self.readings:
            return None
        return float(statistics.median(r.resistance_ohm for r in self.readings))

    @property
    def spread_ohm(self) -> float | None:
        """Return max-min repeat spread, or None when no reading succeeded."""
        if not self.readings:
            return None
        values = [reading.resistance_ohm for reading in self.readings]
        return max(values) - min(values)


@dataclass
class BatchSession:
    """Contain one complete or partial batch run and its solved estimates."""

    target_id: str
    started_at: datetime
    finished_at: datetime
    repeat_count: int
    settle_seconds: float
    cooldown_seconds: float
    pair_samples: dict[str, PairSamples]
    estimates_ohm: dict[str, float]
    warnings: list[str]
    calibration_applied: bool = False
    calibration_profile_path: str | None = None
    cancelled: bool = False

    @property
    def complete(self) -> bool:
        """Report whether all 48 port-path residual estimates were solved."""
        return len(self.estimates_ohm) == PORTS_PER_BANK * 2 and not self.cancelled


def build_measurement_plan() -> list[PairSpec]:
    """Build 48 independent equations plus three cross-bank validation pairs.

    Each 24-port bank contributes 23 adjacent equations and one X0-X2 anchor.
    Adjacent equations alone are underdetermined by one variable; the anchor
    makes the system uniquely solvable without assuming equal path residuals.
    """
    plan: list[PairSpec] = []
    for bank in ("S1", "S2"):
        for x in range(PORTS_PER_BANK - 1):
            plan.append(
                PairSpec(
                    MatrixPort(bank, x),
                    MatrixPort(bank, x + 1),
                    "equation",
                )
            )
        plan.append(
            PairSpec(MatrixPort(bank, 0), MatrixPort(bank, 2), "anchor")
        )

    for x in (0, 12, 23):
        plan.append(
            PairSpec(MatrixPort("S1", x), MatrixPort("S2", x), "validation")
        )
    return plan


def parse_measurement_payload(payload: str) -> CalibrationReading | None:
    """Decode only successful MEASURE payloads while preserving errors raw."""
    match = MEASURE_OK_PATTERN.fullmatch(payload.strip())
    if match is None:
        return None
    return CalibrationReading(
        resistance_ohm=float(match.group("resistance")),
        raw_value=int(match.group("raw")),
        range_code=int(match.group("range")),
        payload=payload,
        received_at=datetime.now().astimezone().isoformat(timespec="seconds"),
    )


def _pair_key(bank: str, first_x: int, second_x: int) -> str:
    return PairSpec(
        MatrixPort(bank, first_x), MatrixPort(bank, second_x), "equation"
    ).key


def solve_path_resistances(
    pair_samples: dict[str, PairSamples],
) -> tuple[dict[str, float], list[str]]:
    """Solve one externally shorted port-path residual from pair sums.

    For each bank, M01=R0+R1, M12=R1+R2 and M02=R0+R2 give
    R0=(M01-M12+M02)/2. Remaining values propagate along adjacent equations.
    Missing measurements leave that bank unsolved instead of inventing values.
    """
    estimates: dict[str, float] = {}
    warnings: list[str] = []

    for bank in ("S1", "S2"):
        required_keys = [
            _pair_key(bank, x, x + 1) for x in range(PORTS_PER_BANK - 1)
        ] + [_pair_key(bank, 0, 2)]
        missing = [
            key
            for key in required_keys
            if key not in pair_samples
            or pair_samples[key].representative_ohm is None
        ]
        if missing:
            warnings.append(
                f"{bank} 缺少 {len(missing)} 个有效方程，未生成该组端口路径阻值。"
            )
            continue

        m01 = pair_samples[_pair_key(bank, 0, 1)].representative_ohm
        m12 = pair_samples[_pair_key(bank, 1, 2)].representative_ohm
        m02 = pair_samples[_pair_key(bank, 0, 2)].representative_ohm
        assert m01 is not None and m12 is not None and m02 is not None

        bank_values = [(m01 - m12 + m02) / 2.0]
        for x in range(PORTS_PER_BANK - 1):
            pair_value = pair_samples[
                _pair_key(bank, x, x + 1)
            ].representative_ohm
            assert pair_value is not None
            bank_values.append(pair_value - bank_values[-1])

        for x, value in enumerate(bank_values):
            label = MatrixPort(bank, x).label
            estimates[label] = value
            if value < 0.0:
                warnings.append(f"{label} 解得负阻值 {value:.3f} ohm，请复测。")

    for x in (0, 12, 23):
        spec = PairSpec(MatrixPort("S1", x), MatrixPort("S2", x), "validation")
        sample = pair_samples.get(spec.key)
        measured = sample.representative_ohm if sample is not None else None
        left = estimates.get(spec.positive.label)
        right = estimates.get(spec.negative.label)
        if measured is None or left is None or right is None:
            warnings.append(f"{spec.label} 缺少跨芯片校验数据。")
            continue
        deviation = measured - (left + right)
        if abs(deviation) > 10.0:
            warnings.append(
                f"{spec.label} 跨芯片校验偏差 {deviation:+.3f} ohm 超过 10 ohm。"
            )

    for samples in pair_samples.values():
        spread = samples.spread_ohm
        if spread is not None and spread > 10.0:
            warnings.append(
                f"{samples.spec.label} 三轮极差 {spread:.3f} ohm 超过 10 ohm。"
            )
        if samples.errors:
            warnings.append(
                f"{samples.spec.label} 有 {len(samples.errors)} 次失败响应。"
            )

    return estimates, warnings


class BatchCalibrationController:
    """Run external-short calibration and correlate asynchronous RESULT frames."""

    def __init__(
        self,
        send_request: SendRequest,
        publish_event: PublishEvent,
        *,
        report_root: Path = REPORT_ROOT,
        repeat_count: int = DEFAULT_REPEAT_COUNT,
        settle_seconds: float = DEFAULT_SETTLE_SECONDS,
        cooldown_seconds: float = DEFAULT_COOLDOWN_SECONDS,
        response_timeout_seconds: float = DEFAULT_RESPONSE_TIMEOUT_SECONDS,
        report_writer: ReportWriter | None = None,
    ) -> None:
        self._send_request = send_request
        self._publish_event = publish_event
        self._report_root = Path(report_root)
        self._repeat_count = repeat_count
        self._settle_seconds = settle_seconds
        self._cooldown_seconds = cooldown_seconds
        self._response_timeout_seconds = response_timeout_seconds
        self._report_writer = report_writer or write_report_files
        self._lock = threading.Lock()
        self._pending: dict[str, tuple[str, queue.Queue[str]]] = {}
        self._thread: threading.Thread | None = None
        self._cancel = threading.Event()

    @property
    def running(self) -> bool:
        """Return whether a batch worker currently owns the measurement path."""
        with self._lock:
            return self._thread is not None and self._thread.is_alive()

    def start(self, target_id: str) -> bool:
        """Start one batch for an online target; reject overlapping sessions."""
        ensure_report_dependency()
        with self._lock:
            if self._thread is not None and self._thread.is_alive():
                return False
            self._cancel.clear()
            thread = threading.Thread(
                target=self._run,
                args=(target_id,),
                daemon=True,
                name="cable-batch-calibration",
            )
            self._thread = thread
            thread.start()
        return True

    def cancel(self) -> None:
        """Request a cooperative stop without interrupting an in-flight command."""
        self._cancel.set()

    def feed_result(self, target_id: str, request_id: str, payload: str) -> bool:
        """Deliver a correlated ESP RESULT to the waiting worker when applicable."""
        with self._lock:
            pending = self._pending.get(request_id)
        if pending is None:
            return False
        expected_target_id, destination = pending
        if target_id != expected_target_id:
            return False
        try:
            destination.put_nowait(payload)
        except queue.Full:
            return False
        return True

    def _publish(self, event_type: str, message: object) -> None:
        self._publish_event(event_type, message)

    def _run(self, target_id: str) -> None:
        started_at = datetime.now().astimezone()
        terminal_event: tuple[str, object] | None = None
        plan = build_measurement_plan()
        samples = {spec.key: PairSamples(spec) for spec in plan}
        total_commands = len(plan) * self._repeat_count
        estimated_minutes = total_commands * (
            self._settle_seconds + self._cooldown_seconds
        ) / 60.0
        sequence = 0
        self._publish(
            "batch_started",
            f"目标 {target_id}，共 {total_commands} 次测量，预计约 "
            f"{estimated_minutes:.1f} 分钟。",
        )

        try:
            for round_index in range(self._repeat_count):
                for spec in plan:
                    if self._cancel.is_set():
                        raise InterruptedError("批量校准已取消")
                    sequence += 1
                    self._publish(
                        "batch_progress",
                        f"{sequence}/{total_commands} 第 {round_index + 1} 轮 "
                        f"{spec.label}",
                    )
                    request_base = (
                        f"BCL-{started_at.strftime('%H%M%S')}-{sequence:03d}"
                    )
                    connect_payload = self._request_one(
                        target_id, f"{request_base}-C", spec.command
                    )
                    if not connect_payload.startswith("OK CONNECT "):
                        samples[spec.key].errors.append(connect_payload)
                        if self._cancel.wait(self._cooldown_seconds):
                            raise InterruptedError("批量校准已取消")
                        continue
                    if self._cancel.wait(self._settle_seconds):
                        raise InterruptedError("批量校准已取消")
                    payload = self._request_one(
                        target_id, f"{request_base}-M", "MEASURE"
                    )
                    reading = parse_measurement_payload(payload)
                    if reading is None:
                        samples[spec.key].errors.append(payload)
                    else:
                        samples[spec.key].readings.append(reading)
                    if self._cancel.wait(self._cooldown_seconds):
                        raise InterruptedError("批量校准已取消")

            estimates, warnings = solve_path_resistances(samples)
            session = BatchSession(
                target_id=target_id,
                started_at=started_at,
                finished_at=datetime.now().astimezone(),
                repeat_count=self._repeat_count,
                settle_seconds=self._settle_seconds,
                cooldown_seconds=self._cooldown_seconds,
                pair_samples=samples,
                estimates_ohm=estimates,
                warnings=warnings,
            )
            docx_path, json_path = self._report_writer(session, self._report_root)
            terminal_event = (
                "batch_complete",
                {
                    "docx": str(docx_path),
                    "json": str(json_path),
                    "target_id": target_id,
                    "solved": len(estimates),
                    "warnings": len(warnings),
                    "calibration_applied": session.calibration_applied,
                    "calibration_profile": session.calibration_profile_path,
                },
            )
        except InterruptedError as error:
            terminal_event = ("batch_error", str(error))
        except Exception as error:  # Keep the GUI alive and surface the exact failure.
            terminal_event = ("batch_error", f"批量校准失败：{error}")
        finally:
            reset_id = f"BCL-{started_at.strftime('%H%M%S')}-RESET"
            try:
                reset_payload = self._request_one(target_id, reset_id, "RESET")
                if reset_payload != "OK RESET":
                    self._publish(
                        "batch_reset_warning",
                        f"批量校准结束复位失败：{reset_payload}",
                    )
            except Exception as error:
                self._publish(
                    "batch_reset_warning",
                    f"批量校准结束复位异常：{error}",
                )
            with self._lock:
                self._pending.clear()
                self._thread = None
            if terminal_event is not None:
                self._publish(*terminal_event)

    def _request_one(self, target_id: str, request_id: str, command: str) -> str:
        response_queue: queue.Queue[str] = queue.Queue(maxsize=1)
        with self._lock:
            self._pending[request_id] = (target_id, response_queue)
        try:
            self._publish("sent", f"GUI -> {target_id} {request_id} {command}")
            acknowledgement = self._send_request(target_id, request_id, command)
            self._publish("received", acknowledgement)
            if not acknowledgement.startswith("OK FORWARDED "):
                return acknowledgement
            try:
                return response_queue.get(timeout=self._response_timeout_seconds)
            except queue.Empty:
                return f"ERR BATCH RESULT_TIMEOUT request_id={request_id}"
        finally:
            with self._lock:
                self._pending.pop(request_id, None)


def _next_report_paths(
    report_root: Path, target_id: str, now: datetime
) -> tuple[str, Path, Path]:
    month_dir = report_root / now.strftime("%Y-%m")
    month_dir.mkdir(parents=True, exist_ok=True)
    safe_target = re.sub(r"[^A-Za-z0-9_-]+", "_", target_id).strip("_") or "DEVICE"
    date_code = now.strftime("%Y%m%d")
    prefix = f"外部短接批量校准_{safe_target}_{date_code}_"
    existing_numbers: list[int] = []
    for existing in month_dir.glob(f"{prefix}*.docx"):
        match = re.search(r"_(\d{3})\.docx$", existing.name)
        if match is not None:
            existing_numbers.append(int(match.group(1)))
    report_number = max(existing_numbers, default=0) + 1
    report_id = f"{date_code}-{report_number:03d}"
    stem = f"{prefix}{report_number:03d}"
    return report_id, month_dir / f"{stem}.docx", month_dir / f"{stem}.json"


def _session_to_json(session: BatchSession, report_id: str) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "report_id": report_id,
        "target_id": session.target_id,
        "started_at": session.started_at.isoformat(timespec="seconds"),
        "finished_at": session.finished_at.isoformat(timespec="seconds"),
        "repeat_count": session.repeat_count,
        "settle_seconds": session.settle_seconds,
        "cooldown_seconds": session.cooldown_seconds,
        "scope": "externally shorted Kelvin port-path residual estimate",
        "calibration_applied": session.calibration_applied,
        "calibration_profile_path": session.calibration_profile_path,
        "complete": session.complete,
        "estimates_ohm": session.estimates_ohm,
        "warnings": session.warnings,
        "pairs": [
            {
                "key": samples.spec.key,
                "positive": samples.spec.positive.label,
                "negative": samples.spec.negative.label,
                "purpose": samples.spec.purpose,
                "representative_ohm": samples.representative_ohm,
                "spread_ohm": samples.spread_ohm,
                "readings": [
                    {
                        "resistance_ohm": reading.resistance_ohm,
                        "raw_value": reading.raw_value,
                        "range_code": reading.range_code,
                        "payload": reading.payload,
                        "received_at": reading.received_at,
                    }
                    for reading in samples.readings
                ],
                "errors": samples.errors,
            }
            for samples in session.pair_samples.values()
        ],
    }


def write_report_files(session: BatchSession, report_root: Path) -> tuple[Path, Path]:
    """Write a dated numbered DOCX and an unapplied JSON calibration candidate."""
    report_id, docx_path, json_path = _next_report_paths(
        Path(report_root), session.target_id, session.finished_at
    )
    payload = _session_to_json(session, report_id)

    json_temp = json_path.with_suffix(".tmp.json")
    json_temp.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    json_temp.replace(json_path)

    docx_temp = docx_path.with_suffix(".tmp.docx")
    _build_docx(session, report_id, docx_temp)
    docx_temp.replace(docx_path)
    return docx_path, json_path


def _build_docx(session: BatchSession, report_id: str, output_path: Path) -> None:
    """Build the compact-reference Word report with fixed table geometry."""
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as error:
        requirements = REQUIREMENTS
        raise RuntimeError(
            "生成Word报告需要python-docx，请执行："
            f'"{sys.executable}" -m pip install -r "{requirements}"'
        ) from error

    blue = RGBColor(0x2E, 0x74, 0xB5)
    dark_blue = RGBColor(0x1F, 0x4D, 0x78)
    muted = RGBColor(0x66, 0x66, 0x66)

    def set_font(run: Any, size: float, *, bold: bool = False, color: Any = None) -> None:
        run.font.name = "Calibri"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(size)
        run.bold = bold
        if color is not None:
            run.font.color.rgb = color

    def set_cell_margins(cell: Any) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for edge, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
            node = tc_mar.find(qn(f"w:{edge}"))
            if node is None:
                node = OxmlElement(f"w:{edge}")
                tc_mar.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    def shade_cell(cell: Any, fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = tc_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            tc_pr.append(shading)
        shading.set(qn("w:fill"), fill)

    def set_table_geometry(table: Any, widths: list[int]) -> None:
        if sum(widths) != 9360:
            raise ValueError("DOCX table widths must sum to 9360 DXA")
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        table_pr = table._tbl.tblPr
        table_width = table_pr.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            table_pr.append(table_width)
        table_width.set(qn("w:w"), "9360")
        table_width.set(qn("w:type"), "dxa")
        indent = table_pr.find(qn("w:tblInd"))
        if indent is None:
            indent = OxmlElement("w:tblInd")
            table_pr.append(indent)
        indent.set(qn("w:w"), "120")
        indent.set(qn("w:type"), "dxa")

        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for width in widths:
            grid_column = OxmlElement("w:gridCol")
            grid_column.set(qn("w:w"), str(width))
            grid.append(grid_column)

        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_width = tc_pr.find(qn("w:tcW"))
                if tc_width is None:
                    tc_width = OxmlElement("w:tcW")
                    tc_pr.append(tc_width)
                tc_width.set(qn("w:w"), str(width))
                tc_width.set(qn("w:type"), "dxa")
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)

    def style_table(table: Any, widths: list[int], numeric_columns: set[int]) -> None:
        set_table_geometry(table, widths)
        table.style = "Table Grid"
        header = table.rows[0]
        header_property = header._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_property.append(repeat)
        for column, cell in enumerate(header.cells):
            shade_cell(cell, "E8EEF5")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_font(run, 9, bold=True, color=dark_blue)
        for row in table.rows[1:]:
            for column, cell in enumerate(row.cells):
                for paragraph in cell.paragraphs:
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if column in numeric_columns
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        set_font(run, 8.5)

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in (
        ("Heading 1", 16, blue, 18, 10),
        ("Heading 2", 13, blue, 14, 7),
        ("Heading 3", 12, dark_blue, 10, 5),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(
        header.add_run("ESP32 Cable Tester | External-Short Calibration"),
        9,
        color=muted,
    )
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run(f"Report {report_id} | Page "), 9, color=muted)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    for field_element in (field_begin, instruction, field_end):
        field_run = footer.add_run()
        set_font(field_run, 9, color=muted)
        field_run._r.append(field_element)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("外部短接批量校准报告"), 23, bold=True)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_font(subtitle.add_run("相邻方程求解与跨芯片一致性校验"), 14, color=muted)

    metadata = (
        ("报告编号", report_id),
        ("目标设备", session.target_id),
        ("开始时间", session.started_at.strftime("%Y-%m-%d %H:%M:%S %z")),
        ("结束时间", session.finished_at.strftime("%Y-%m-%d %H:%M:%S %z")),
        (
            "测量设置",
            f"{session.repeat_count}轮，闭合稳定{session.settle_seconds:.1f}秒，"
            f"命令间冷却{session.cooldown_seconds:.1f}秒",
        ),
        (
            "求解状态",
            f"{len(session.estimates_ohm)}/48路，"
            f"上位机显示校准{'已启用' if session.calibration_applied else '未启用'}",
        ),
    )
    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        set_font(paragraph.add_run(f"{label}: "), 10.5, bold=True)
        set_font(paragraph.add_run(value), 10.5)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note.paragraph_format.space_after = Pt(10)
    set_font(note.add_run("范围说明: "), 10.5, bold=True, color=dark_blue)
    set_font(
        note.add_run(
            "本报告估算48个端口在普通四线测量路径中的残余阻值。测试夹具必须在外部"
            "可靠短接全部端口；结果包含端口、连接器和外部短接夹具的影响，不代表"
            "CH446的Y0-Y3交叉点导通电阻。"
            + (
                "结果已按设备ID写入上位机显示校准配置，未写入ESP32 NVS。"
                if session.calibration_applied
                else "结果仅保存为候选，本次未应用归零参数。"
            )
        ),
        10.5,
    )

    document.add_heading("1 计算方法", level=1)
    document.add_paragraph(
        "外部短接后，普通四线测量得到两个端口路径残余之和: M(i,j)=R(i)+R(j)。"
        "每组先测X0-X1、X1-X2等23个相邻对，再测X0-X2作为独立锚点。"
        "由R0=(M01-M12+M02)/2得到首项，随后按R(i+1)=M(i,i+1)-R(i)递推。"
    )
    document.add_paragraph(
        "每个方程取多轮成功读数的中位数。S1-S2的X0、X12、X23只用于检查"
        "外部短接夹具和两组解的一致性。"
    )

    document.add_heading("2 端口路径估算结果", level=1)
    result_table = document.add_table(rows=1, cols=4)
    for cell, text in zip(
        result_table.rows[0].cells,
        ("矩阵组", "X通道", "路径残余/ohm", "质量标记"),
    ):
        cell.text = text
    for bank in ("S1", "S2"):
        for x in range(PORTS_PER_BANK):
            label = MatrixPort(bank, x).label
            value = session.estimates_ohm.get(label)
            quality = (
                "未解"
                if value is None
                else (
                    "负值-复测"
                    if value < 0
                    else ("已启用" if session.calibration_applied else "候选")
                )
            )
            cells = result_table.add_row().cells
            cells[0].text = bank
            cells[1].text = str(x)
            cells[2].text = "-" if value is None else f"{value:.3f}"
            cells[3].text = quality
    style_table(result_table, [1100, 1100, 3000, 4160], {0, 1, 2, 3})

    document.add_heading("3 跨芯片校验", level=1)
    validation_table = document.add_table(rows=1, cols=5)
    for cell, text in zip(
        validation_table.rows[0].cells,
        ("测量对", "实测/ohm", "解预测/ohm", "偏差/ohm", "状态"),
    ):
        cell.text = text
    for x in (0, 12, 23):
        spec = PairSpec(MatrixPort("S1", x), MatrixPort("S2", x), "validation")
        measured = session.pair_samples[spec.key].representative_ohm
        left = session.estimates_ohm.get(spec.positive.label)
        right = session.estimates_ohm.get(spec.negative.label)
        predicted = None if left is None or right is None else left + right
        deviation = (
            None if measured is None or predicted is None else measured - predicted
        )
        status = "缺数据" if deviation is None else ("复测" if abs(deviation) > 10 else "通过")
        cells = validation_table.add_row().cells
        cells[0].text = spec.label
        cells[1].text = "-" if measured is None else f"{measured:.3f}"
        cells[2].text = "-" if predicted is None else f"{predicted:.3f}"
        cells[3].text = "-" if deviation is None else f"{deviation:+.3f}"
        cells[4].text = status
    style_table(validation_table, [2600, 1550, 1550, 1550, 2110], {1, 2, 3, 4})

    document.add_heading("4 原始测量汇总", level=1)
    raw_table = document.add_table(rows=1, cols=6)
    for cell, text in zip(
        raw_table.rows[0].cells,
        ("测量对", "用途", "成功读数/ohm", "中位数", "极差", "失败数"),
    ):
        cell.text = text
    for samples in session.pair_samples.values():
        values = ", ".join(f"{r.resistance_ohm:.3f}" for r in samples.readings) or "-"
        median = samples.representative_ohm
        spread = samples.spread_ohm
        cells = raw_table.add_row().cells
        cells[0].text = samples.spec.label
        cells[1].text = samples.spec.purpose
        cells[2].text = values
        cells[3].text = "-" if median is None else f"{median:.3f}"
        cells[4].text = "-" if spread is None else f"{spread:.3f}"
        cells[5].text = str(len(samples.errors))
    style_table(raw_table, [2200, 1100, 2500, 1200, 1100, 1260], {3, 4, 5})

    document.add_heading("5 错误响应原文", level=1)
    error_rows = [
        (samples.spec.label, occurrence, payload)
        for samples in session.pair_samples.values()
        for occurrence, payload in enumerate(samples.errors, start=1)
    ]
    if error_rows:
        error_table = document.add_table(rows=1, cols=3)
        for cell, text in zip(
            error_table.rows[0].cells,
            ("测量对", "序号", "ESP原始错误响应"),
        ):
            cell.text = text
        for pair_label, occurrence, payload in error_rows:
            cells = error_table.add_row().cells
            cells[0].text = pair_label
            cells[1].text = str(occurrence)
            cells[2].text = payload
        style_table(error_table, [2300, 800, 6260], {1})
    else:
        document.add_paragraph("本次批量校准没有收到错误响应。")

    document.add_heading("6 异常与结论", level=1)
    if session.warnings:
        for index, warning in enumerate(session.warnings, start=1):
            paragraph = document.add_paragraph()
            set_font(paragraph.add_run(f"W{index:02d}  "), 10, bold=True, color=dark_blue)
            set_font(paragraph.add_run(warning), 10)
    else:
        document.add_paragraph("未发现缺失方程、负阻值、过大重复极差或跨芯片校验超差。")
    conclusion = document.add_paragraph()
    conclusion.paragraph_format.space_before = Pt(8)
    set_font(conclusion.add_run("结论: "), 11, bold=True, color=dark_blue)
    set_font(
        conclusion.add_run(
            (
                "本次结果已用于该设备的上位机显示校准；正常测量会分别扣除正、负"
                "端口阻值，ESP32原始返回和NVS保持不变。"
                if session.calibration_applied
                else "本次结果可作为后续校准评审输入；软件不会自动扣除这些阻值。"
            )
        ),
        11,
    )

    document.core_properties.title = "外部短接批量校准报告"
    document.core_properties.subject = "普通四线测量端口路径残余阻值求解"
    document.core_properties.author = "Cable Tester GUI"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def build_demo_session() -> BatchSession:
    """Create deterministic synthetic data for tests and DOCX render QA only."""
    plan = build_measurement_plan()
    samples = {spec.key: PairSamples(spec) for spec in plan}
    true_values = {
        MatrixPort(bank, x).label: 108.0 + (4.0 if bank == "S2" else 0.0) + x * 0.45
        for bank in ("S1", "S2")
        for x in range(PORTS_PER_BANK)
    }
    for spec in plan:
        total = true_values[spec.positive.label] + true_values[spec.negative.label]
        for offset in (-0.4, 0.0, 0.4):
            resistance = total + offset
            samples[spec.key].readings.append(
                CalibrationReading(
                    resistance_ohm=resistance,
                    raw_value=round(resistance * 10),
                    range_code=0,
                    payload=(
                        f"OK MEASURE resistance={resistance:.3f} "
                        f"raw={round(resistance * 10)} range=0"
                    ),
                    received_at="2026-08-10T10:00:00+08:00",
                )
            )
    estimates, warnings = solve_path_resistances(samples)
    return BatchSession(
        target_id="ESP1",
        started_at=datetime.fromisoformat("2026-08-10T10:00:00+08:00"),
        finished_at=datetime.fromisoformat("2026-08-10T10:08:00+08:00"),
        repeat_count=3,
        settle_seconds=0.2,
        cooldown_seconds=0.2,
        pair_samples=samples,
        estimates_ohm=estimates,
        warnings=warnings,
    )
