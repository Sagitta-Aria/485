from __future__ import annotations

import json
import queue
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

import batch_calibration as batch
import cable_tester_gui as gui


class BatchCalibrationMathTests(unittest.TestCase):
    """Verify the measurement plan and the independent-equation solver."""

    def test_plan_contains_equations_anchors_and_validations(self) -> None:
        plan = batch.build_measurement_plan()

        self.assertEqual(batch.DEFAULT_SETTLE_SECONDS, 2.0)
        self.assertEqual(batch.DEFAULT_COOLDOWN_SECONDS, 0.2)
        self.assertEqual(len(plan), 51)
        self.assertEqual(plan[0].command, "CONNECT S1 0 S1 1")
        self.assertEqual(plan[1].command, "CONNECT S1 1 S1 2")
        self.assertEqual(
            sum(spec.purpose == "equation" for spec in plan), 46
        )
        self.assertEqual(sum(spec.purpose == "anchor" for spec in plan), 2)
        self.assertEqual(sum(spec.purpose == "validation" for spec in plan), 3)

    def test_missing_docx_error_uses_current_interpreter(self) -> None:
        with mock.patch("builtins.__import__", side_effect=ImportError("missing")):
            with self.assertRaises(RuntimeError) as caught:
                batch.ensure_report_dependency()

        self.assertIn(sys.executable, str(caught.exception))
        self.assertIn("requirements.txt", str(caught.exception))

    def test_solver_recovers_all_demo_path_values(self) -> None:
        session = batch.build_demo_session()

        self.assertTrue(session.complete)
        self.assertEqual(len(session.estimates_ohm), 48)
        self.assertEqual(session.warnings, [])
        self.assertAlmostEqual(session.estimates_ohm["S1_X0"], 108.0, places=6)
        self.assertAlmostEqual(session.estimates_ohm["S2_X23"], 122.35, places=6)

    def test_report_writes_unapplied_calibration_candidate(self) -> None:
        session = batch.build_demo_session()
        first_samples = next(iter(session.pair_samples.values()))
        exact_error = "ERR MEASURE OVERRANGE status=1 frame=010304FFFFFFFF1234"
        first_samples.errors.append(exact_error)
        with tempfile.TemporaryDirectory() as temporary_directory:
            docx_path, json_path = batch.write_report_files(
                session, Path(temporary_directory)
            )

            self.assertTrue(docx_path.exists())
            self.assertTrue(json_path.exists())
            payload = json.loads(json_path.read_text(encoding="utf-8"))
            self.assertFalse(payload["calibration_applied"])
            self.assertEqual(payload["settle_seconds"], 0.2)
            self.assertEqual(
                payload["scope"],
                "externally shorted Kelvin port-path residual estimate",
            )
            self.assertEqual(len(payload["estimates_ohm"]), 48)
            self.assertEqual(payload["pairs"][0]["errors"], [exact_error])

            from docx import Document

            report = Document(docx_path)
            report_text = "\n".join(
                cell.text
                for table in report.tables
                for row in table.rows
                for cell in row.cells
            )
            self.assertIn(exact_error, report_text)

    def test_report_records_gui_profile_activation(self) -> None:
        session = batch.build_demo_session()
        session.calibration_applied = True
        session.calibration_profile_path = (
            "calibration_profiles/ESP1/active.json"
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            docx_path, json_path = batch.write_report_files(
                session, Path(temporary_directory)
            )

            payload = json.loads(json_path.read_text(encoding="utf-8"))
            self.assertTrue(payload["calibration_applied"])
            self.assertEqual(
                payload["calibration_profile_path"],
                "calibration_profiles/ESP1/active.json",
            )

            from docx import Document

            report = Document(docx_path)
            paragraph_text = "\n".join(
                paragraph.text for paragraph in report.paragraphs
            )
            self.assertIn("上位机显示校准已启用", paragraph_text)
            self.assertIn("未写入ESP32 NVS", paragraph_text)


class BatchCalibrationControllerTests(unittest.TestCase):
    """Exercise request/result correlation without real hardware or delays."""

    def test_controller_completes_with_correlated_results(self) -> None:
        events: list[tuple[str, object]] = []
        true_values = {
            batch.MatrixPort(bank, x).label: 100.0 + (5.0 if bank == "S2" else 0.0) + x
            for bank in ("S1", "S2")
            for x in range(batch.PORTS_PER_BANK)
        }

        controller: batch.BatchCalibrationController
        selected_pair: tuple[str, int, str, int] | None = None
        sent_commands: list[str] = []

        def send_request(target_id: str, request_id: str, command: str) -> str:
            nonlocal selected_pair
            sent_commands.append(command)
            if command.startswith("CONNECT "):
                _, first_bank, first_x, second_bank, second_x = command.split()
                selected_pair = (
                    first_bank,
                    int(first_x),
                    second_bank,
                    int(second_x),
                )
                payload = (
                    f"OK CONNECT {first_bank} {first_x} "
                    f"{second_bank} {second_x}"
                )
            elif command == "MEASURE":
                assert selected_pair is not None
                first_bank, first_x, second_bank, second_x = selected_pair
                resistance = (
                    true_values[batch.MatrixPort(first_bank, first_x).label]
                    + true_values[batch.MatrixPort(second_bank, second_x).label]
                )
                payload = (
                    f"OK MEASURE resistance={resistance:.3f} "
                    f"raw={round(resistance * 10)} range=0"
                )
            elif command == "RESET":
                payload = "OK RESET"
            else:
                self.fail(f"unexpected command: {command}")
            self.assertTrue(controller.feed_result(target_id, request_id, payload))
            return f"OK FORWARDED {target_id} {request_id}"

        def write_fake_report(
            session: batch.BatchSession, report_root: Path
        ) -> tuple[Path, Path]:
            self.assertEqual(len(session.estimates_ohm), 48)
            self.assertFalse(session.calibration_applied)
            return report_root / "report.docx", report_root / "report.json"

        controller = batch.BatchCalibrationController(
            send_request,
            lambda event_type, message: events.append((event_type, message)),
            report_root=Path("reports"),
            repeat_count=1,
            settle_seconds=0.0,
            cooldown_seconds=0.0,
            response_timeout_seconds=0.1,
            report_writer=write_fake_report,
        )
        self.assertTrue(controller.start("ESP1"))

        deadline = time.monotonic() + 2.0
        while controller.running and time.monotonic() < deadline:
            time.sleep(0.01)

        self.assertFalse(controller.running)
        completed = [message for event, message in events if event == "batch_complete"]
        errors = [message for event, message in events if event == "batch_error"]
        self.assertEqual(errors, [])
        self.assertEqual(len(completed), 1)
        self.assertEqual(completed[0]["solved"], 48)
        self.assertFalse(completed[0]["calibration_applied"])
        self.assertEqual(sent_commands[:2], ["CONNECT S1 0 S1 1", "MEASURE"])
        self.assertEqual(sent_commands[-1], "RESET")
        self.assertEqual(sent_commands.count("MEASURE"), 51)

    def test_controller_rejects_result_from_wrong_target(self) -> None:
        controller = batch.BatchCalibrationController(
            lambda target_id, request_id, command: "OK FORWARDED ESP1 REQ-1",
            lambda event_type, message: None,
        )
        destination: queue.Queue[str] = queue.Queue(maxsize=1)
        controller._pending["REQ-1"] = ("ESP1", destination)

        self.assertFalse(
            controller.feed_result(
                "ESP2",
                "REQ-1",
                "OK MEASURE resistance=220.000 raw=2200 range=0",
            )
        )
        self.assertTrue(destination.empty())
        self.assertTrue(
            controller.feed_result(
                "ESP1",
                "REQ-1",
                "OK MEASURE resistance=220.000 raw=2200 range=0",
            )
        )

    def test_cancel_resets_before_publishing_terminal_event(self) -> None:
        commands: list[str] = []
        events: list[tuple[str, object]] = []
        controller: batch.BatchCalibrationController

        def send_request(target_id: str, request_id: str, command: str) -> str:
            commands.append(command)
            if command.startswith("CONNECT "):
                payload = "OK CONNECT S1 0 S1 1"
                controller.cancel()
            elif command == "RESET":
                payload = "OK RESET"
            else:
                self.fail(f"unexpected command after cancellation: {command}")
            self.assertTrue(controller.feed_result(target_id, request_id, payload))
            return f"OK FORWARDED {target_id} {request_id}"

        controller = batch.BatchCalibrationController(
            send_request,
            lambda event_type, message: events.append((event_type, message)),
            settle_seconds=0.0,
            cooldown_seconds=0.0,
            response_timeout_seconds=0.1,
        )
        self.assertTrue(controller.start("ESP1"))

        deadline = time.monotonic() + 1.0
        while controller.running and time.monotonic() < deadline:
            time.sleep(0.01)

        self.assertEqual(commands, ["CONNECT S1 0 S1 1", "RESET"])
        terminal = [event for event in events if event[0] == "batch_error"]
        self.assertEqual(len(terminal), 1)
        self.assertEqual(events[-1][0], "batch_error")

    def test_gui_controls_one_crosspoint_without_reset(self) -> None:
        app = gui.CableTesterApp.__new__(gui.CableTesterApp)
        app.switch_bank = mock.Mock()
        app.switch_bank.get.return_value = "S2"
        app.switch_x = mock.Mock()
        app.switch_x.get.return_value = "11"
        app.switch_bus = mock.Mock()
        app.switch_bus.get.return_value = "Y4"
        app._send = mock.Mock()

        app._set_crosspoint(True)
        app._set_crosspoint(False)

        self.assertEqual(
            app._send.call_args_list,
            [
                mock.call("SWITCH S2 11 Y4 ON"),
                mock.call("SWITCH S2 11 Y4 OFF"),
            ],
        )

    def test_gui_router_publishes_structured_result_frame(self) -> None:
        events: queue.Queue[tuple[str, object]] = queue.Queue()
        state = gui.RouterState(events)
        target = mock.Mock(peer_id="ESP1")

        error = state.route_result(
            target,
            "ESP1",
            gui.CONTROLLER_ID,
            "BST-100000-001",
            "OK MEASURE resistance=220.000 raw=2200 range=0",
        )

        self.assertIsNone(error)
        self.assertEqual(
            events.get_nowait(),
            (
                "result_frame",
                (
                    "ESP1",
                    "BST-100000-001",
                    "OK MEASURE resistance=220.000 raw=2200 range=0",
                ),
            ),
        )
        self.assertTrue(events.empty())


if __name__ == "__main__":
    unittest.main()
